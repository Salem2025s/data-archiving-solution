"""Génère le document Word des annexes du Bloc 2, images embarquées.

Annexes A→H (analyse du besoin, plan, dashboard, tests statistiques,
visualisations, recommandations, formation, documentation technique).
Captures : dossier « Bloc 2 » en priorité, repli sur « Bloc 1 » pour les
pages réutilisées (Domaines, Console SQL).

Usage : python scripts/build_bloc2_annexes.py
"""
from __future__ import annotations

import glob
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMG_DIRS = [
    ROOT / "docs" / "certification" / "Bloc 2",
    ROOT / "docs" / "certification" / "Bloc 1",
]
OUT = ROOT / "docs" / "certification" / "BLOC2_Annexes.docx"

ACCENT = RGBColor(0x25, 0x63, 0xEB)
GREY = RGBColor(0x4B, 0x55, 0x63)
CODE_BG = "F1F3F8"
IMG_W = Cm(16.0)


def find_img(prefix: str) -> str | None:
    for d in IMG_DIRS:
        hits = [m for m in glob.glob(str(d / f"{prefix}*"))
                if m.lower().endswith((".png", ".jpg", ".jpeg"))]
        if hits:
            return sorted(hits)[0]
    return None


def style_base(doc: Document) -> None:
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10.5)
    for h, sz, col in (("Heading 1", 16, ACCENT), ("Heading 2", 13, ACCENT),
                       ("Heading 3", 11.5, GREY)):
        s = doc.styles[h]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.color.rgb = col


def add_bold(p, text: str) -> None:
    import re
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        if i % 2 == 1:
            r.bold = True


def para(doc, text: str) -> None:
    add_bold(doc.add_paragraph(), text)


def code(doc, text: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(tcpr.makeelement(qn("w:shd"),
                {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): CODE_BG}))
    cell.paragraphs[0].text = ""
    for idx, ln in enumerate(text.split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(ln or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    tblpr = t._tbl.tblPr
    b = tblpr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right"):
        b.append(tblpr.makeelement(qn(f"w:{edge}"),
                 {qn("w:val"): "single", qn("w:sz"): "4", qn("w:color"): "D0D7E5"}))
    tblpr.append(b)


def table(doc, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.paragraphs[0].text = ""
        add_bold(c.paragraphs[0], h)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].paragraphs[0].text = ""
            add_bold(cells[j].paragraphs[0], val)


def image(doc, prefix: str, caption: str) -> None:
    path = find_img(prefix)
    if not path:
        para(doc, f"_[capture {prefix} introuvable]_")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=IMG_W)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


# ===========================================================================
def build() -> None:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    style_base(doc)

    doc.add_heading("BLOC 2 — Annexes", level=1)
    para(doc, "**Titre RNCP 39586 · Bloc 2 — Analyser, organiser et valoriser des données**")
    para(doc, "Candidat : HAOUARI Salem · Date : 27/07/2026 · Annexes au dossier écrit (hors décompte de pages)")
    para(doc, "Ces annexes apportent les **preuves d'analyse** décrites dans le dossier. "
              "Les chiffres proviennent d'exécutions réelles sur l'entrepôt PostgreSQL "
              "pfe_data_ia (run 9).")

    # ---- Annexe A ----
    doc.add_heading("Annexe A — Analyse du besoin", level=2)
    para(doc, "Synthèse du cadrage de l'analyse à partir du besoin exprimé par le commanditaire (DSI).")
    table(doc, ["Dimension", "Contenu"], [
        ["Enjeux / problématique", "Cartographier, chiffrer et recommander l'archivage d'un patrimoine ERP sans risque"],
        ["Contexte", "ERP PeopleSoft en forte croissance, sans cartographie ni chiffrage à jour"],
        ["Environnement", "Source Oracle EP92U038 (lecture seule) → entrepôt PostgreSQL médaillon ; Python, SQL, dashboard, Excel"],
        ["Contrainte technique", "VPN, aucune écriture sur la source, noms de tables opaques, >160 k actifs"],
        ["Contrainte réglementaire", "RGPD ; durées légales de rétention"],
        ["Contrainte coût", "Prototype : outils open-source imposés (PostgreSQL, Python)"],
        ["Contrainte délai / produit", "Analyses rejouables (run_id) ; solution générique (base substituable, vendable)"],
    ])

    # ---- Annexe B ----
    doc.add_heading("Annexe B — Plan d'analyse", level=2)
    para(doc, "Traduction de la problématique métier en problème numérique, en trois axes.")
    table(doc, ["Axe", "Métriques", "Traduction numérique"], [
        ["1 · Cartographie par domaine", "Nb actifs/domaine, part, bandes de confiance, taux de revue", "Classification supervisée multiclasse (7 domaines)"],
        ["2 · Coûts & archivage", "Volume (Mo)/domaine, coût chaud, coût archive, économie €/an", "Agrégation de coûts pondérés par la volumétrie"],
        ["3 · Valorisation (ROI, N-1)", "Score ROI, projection pluriannuelle, comparaison N-1", "Scoring + simulation Monte-Carlo"],
    ])
    para(doc, "Toutes les métriques s'appuient sur des tables **déjà produites et versionnées** de la "
              "couche `serving` : le plan est directement exécutable.")

    # ---- Annexe C ----
    doc.add_heading("Annexe C — Requêtes et restitution en dashboard", level=2)
    para(doc, "**C.1 — Vue d'ensemble du patrimoine.** Cartographie des 167 260 actifs : KPIs, "
              "répartition par domaine, volume stocké et synthèse des recommandations d'archivage.")
    image(doc, "Dashboard", "Capture — Page « Vue d'ensemble » : 167 260 actifs, 7 domaines, économie normalisée (C2.1.3 / C2.2.1).")
    para(doc, "**C.2 — Profil par domaine métier.** Volumétrie, score d'archivage et qualité des "
              "données, domaine par domaine.")
    image(doc, "#9 ", "Capture — Page « Domaines » : profil détaillé par objet métier (C2.1.3).")
    para(doc, "**C.3 — Console SQL en lecture seule.** L'analyste exécute ses propres requêtes "
              "sur la base, sans risque d'écriture. Exemple — répartition et volume par domaine :")
    code(doc,
         "SELECT p.business_domain_predicted AS domaine,\n"
         "       COUNT(*)                     AS nb_actifs,\n"
         "       ROUND(SUM(inv.size_mb)::numeric, 1) AS volume_mo\n"
         "FROM serving.asset_business_domain_prediction p\n"
         "JOIN serving.mv_asset_inventory inv\n"
         "  ON inv.asset_id = p.asset_id AND inv.run_id = p.run_id\n"
         "WHERE p.run_id = 9\n"
         "GROUP BY p.business_domain_predicted\n"
         "ORDER BY nb_actifs DESC;")
    image(doc, "#2 ", "Capture — Console SQL : interrogation directe de la couche serving (C2.1.3).")

    # ---- Annexe D ----
    doc.add_heading("Annexe D — Tests statistiques (C2.1.4)", level=2)
    para(doc, "Quatre tests d'hypothèses sur les données réelles du run 9 (seuil α = 0,05), "
              "reproductibles via `scripts/bloc2_hypothesis_tests.py`. La **taille d'effet** est "
              "reportée systématiquement : avec n = 167 260, la seule p-value ne suffit pas à juger "
              "de l'importance d'un résultat.")
    table(doc, ["Hypothèse (H1)", "Test", "Résultat", "Décision"], [
        ["Le nb de colonnes diffère selon le domaine", "Kruskal-Wallis (7 grp, n=60 284)", "H=12 783 · **ε²=0,21 (fort)**", "p<0,001 → rejet de H0"],
        ["Taille corrélée au nb de colonnes", "Spearman (n=15 854)", "**rho=0,368 (modérée)**", "p<0,001 → rejet de H0"],
        ["Description absente selon le domaine", "χ² d'indépendance (ddl=6)", "χ²=3 681 · **V=0,148 (faible)**", "p<0,001 → rejet de H0"],
        ["Confiance keyword ≠ confiance ML", "Mann-Whitney U", "**effet=−0,67 (fort)** · méd. 0,83 vs 0,68", "p<0,001 → rejet de H0"],
    ])
    para(doc, "**Interprétations clés.** Les tables Finance sont structurellement plus riches "
              "(médiane 17 colonnes vs 4 pour « Other ») → oriente la modélisation. La documentation "
              "est inégale selon les domaines → cible la gouvernance. Les deux passes de "
              "classification ne sont pas interchangeables → justifie l'architecture keyword + ML.")
    para(doc, "**Limite assumée.** Une 5ᵉ hypothèse — « le score d'archivabilité diffère selon le "
              "domaine » — a été **écartée après examen** : en périmètre Oracle-only, ce score est "
              "quasi-constant (90 % des actifs à 45,0), faute de signal de purge/accès. Un test "
              "l'aurait déclaré « significatif » à cause du grand n, alors que les médianes sont "
              "identiques — ce serait trompeur.")

    # ---- Annexe E ----
    doc.add_heading("Annexe E — Visualisations et accessibilité", level=2)
    para(doc, "**E.1 — Analyse des coûts et projection ROI.** Répartition du coût par domaine, "
              "**diagramme de Pareto** (identifie les domaines les plus coûteux), et projection ROI "
              "sur 5 ans avec intervalle Monte-Carlo P10–P90.")
    image(doc, "Annexe C Cout", "Capture — Page « Coûts & ROI » : Pareto des coûts et projection ROI (C2.2.1 / C2.2.2).")
    para(doc, "**E.2 — Incertitude du modèle rendue visible.** La page « Test du modèle » affiche la "
              "**distribution des probabilités par domaine** : le jury voit non seulement le label "
              "gagnant mais aussi la confiance calibrée et la répartition sur les 7 classes.")
    image(doc, "test le model", "Capture — « Test du modèle » : prédiction RH à 92,3 % et distribution des probabilités (C2.2.1).")
    para(doc, "**E.3 — Note d'accessibilité (situations de handicap).** Les choix visuels prennent "
              "en compte les déficiences de la vision des couleurs :")
    table(doc, ["Principe", "Mise en œuvre"], [
        ["La couleur n'est jamais seule", "Chaque catégorie porte aussi un libellé et une position (barres triées, légende)"],
        ["Échelles monochromes", "Dégradés d'une seule teinte pour les intensités (pas de rouge-vert)"],
        ["Contraste élevé", "Texte sombre (#0F172A) sur fond clair, conforme WCAG"],
        ["Encodage par forme/longueur", "Les barres encodent la donnée par leur longueur, indépendamment de la couleur"],
    ])

    # ---- Annexe F ----
    doc.add_heading("Annexe F — Recommandations (fiche de synthèse)", level=2)
    para(doc, "Recommandations priorisées par **effet de levier**, destinées au décideur.")
    table(doc, ["#", "Recommandation", "Constat chiffré"], [
        ["1", "Prioriser la gouvernance sur Finance & Contrôle", "45,1 % du patrimoine, tables les plus complexes (méd. 17 colonnes)"],
        ["2", "Traiter le taux de revue comme un chantier qualité", "27,1 % des actifs `review_required`, 25,6 % en confiance low"],
        ["3", "Cibler la documentation sur les domaines sous-documentés", "H3 : l'absence de description dépend du domaine"],
        ["4", "Sécuriser le modèle de coût avant d'industrialiser le ROI", "Tarifs Azure réels (0,2120 vs 0,0216 $/Go/an, écart ×9,8)"],
    ])
    para(doc, "**Synthèse.** Le patrimoine est concentré (Finance ≈ moitié) et inégalement documenté. "
              "Le gain le plus rapide vient de l'effort ciblé sur Finance et les cas incertains, pas "
              "d'un traitement uniforme. Le potentiel d'archivage est réel, mais sa valorisation en "
              "ROI doit attendre un modèle de coût sourcé et un historique de plusieurs runs.")

    # ---- Annexe G ----
    doc.add_heading("Annexe G — Support de formation", level=2)
    para(doc, "Le support de formation destiné aux analystes (non techniciens) est fourni en pièce "
              "jointe : `BLOC2_support_formation.docx`. Il comporte six modules courts — à quoi sert "
              "l'outil, visite guidée du dashboard, lire un KPI sans se tromper, le réflexe qualité, "
              "exporter/réutiliser, bonnes pratiques.")
    para(doc, "**Message pédagogique central :** une confiance *low* ou un marqueur `review_required` "
              "n'est pas un défaut du modèle mais une **alerte utile** — on ne décide jamais "
              "d'archiver sur une classification non vérifiée.")

    # ---- Annexe H ----
    doc.add_heading("Annexe H — Documentation technique (dictionnaire des indicateurs)", level=2)
    para(doc, "Chaque indicateur est adossé à une définition fonctionnelle, une formule technique "
              "et une source, garantissant compréhension et reproductibilité.")
    table(doc, ["Indicateur", "Définition fonctionnelle", "Source / formule"], [
        ["Répartition par domaine", "Poids de chaque objet métier", "`asset_business_domain_prediction`, COUNT par domaine"],
        ["Bande de confiance", "Fiabilité de la classification (high/medium/low)", "`confidence` seuillée (≥0,85 / ≥0,60 / <0,60)"],
        ["Taux de revue", "Part des actifs à vérifier manuellement", "`review_required` = true"],
        ["Volume par domaine", "Stockage occupé par métier", "SUM(`size_mb`) par domaine (`mv_asset_inventory`)"],
        ["Économie annuelle potentielle", "Gain si archivage des candidats", "SUM(size candidats)×(coût_chaud−coût_archive)"],
        ["Score de ROI", "Retour attendu de l'archivage", "`roi_score` (`features_asset`)"],
    ])
    para(doc, "**Sources** : Oracle PeopleSoft EP92U038 (métadonnées, périmètre 167 260 actifs). "
              "**Méthodes** documentées dans `docs/B_…` à `docs/E4_…` et `scripts/bloc2_hypothesis_tests.py`. "
              "**Reproductibilité** assurée par le versionnement `run_id` et les scripts Git.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK -> {OUT}  ({OUT.stat().st_size:,} octets)")


if __name__ == "__main__":
    build()
