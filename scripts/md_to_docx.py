"""Convertit un dossier Markdown (docs/certification/*.md) en fichier Word (.docx).

Gère : titres (#..####), paragraphes, listes à puces, tableaux Markdown,
blocs de code / diagrammes ASCII (```), citations (>), et gras **...** inline.

Usage :
    python scripts/md_to_docx.py docs/certification/BLOC1_Collecter_Transformer_Securiser.md
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x25, 0x63, 0xEB)      # bleu corporate
GREY   = RGBColor(0x4B, 0x55, 0x63)
CODE_BG = "F1F3F8"


def _shade(cell, hex_color: str) -> None:
    """Applique une couleur de fond à une cellule de tableau."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color,
    })
    tc_pr.append(shd)


def _add_runs_with_bold(paragraph, text: str) -> None:
    """Ajoute du texte en gérant le gras **...** inline."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:  # segments capturés = en gras
            run.bold = True


def _style_base(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for h, sz, color in (("Heading 1", 17, ACCENT), ("Heading 2", 14, ACCENT),
                         ("Heading 3", 12, GREY), ("Heading 4", 11, GREY)):
        st = doc.styles[h]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.color.rgb = color


def _add_code_block(doc: Document, lines: list[str]) -> None:
    """Bloc monospace à fond gris (code SQL / diagrammes ASCII)."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade(cell, CODE_BG)
    cell.paragraphs[0].text = ""
    for idx, ln in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(ln if ln else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    # bordure fine
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right"):
        borders.append(tbl_pr.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "4",
            qn("w:color"): "D0D7E5",
        }))
    tbl_pr.append(borders)


def _add_md_table(doc: Document, rows: list[str]) -> None:
    """Convertit un tableau Markdown (lignes avec |) en tableau Word."""
    def cells(line: str) -> list[str]:
        return [c.strip() for c in line.strip().strip("|").split("|")]

    header = cells(rows[0])
    body = [cells(r) for r in rows[2:]]  # rows[1] = séparateur ---
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(header):
        c = table.rows[0].cells[j]
        c.paragraphs[0].text = ""
        _add_runs_with_bold(c.paragraphs[0], htext)
        for run in c.paragraphs[0].runs:
            run.bold = True
    for row in body:
        wr = table.add_row().cells
        for j in range(ncols):
            txt = row[j] if j < len(row) else ""
            wr[j].paragraphs[0].text = ""
            _add_runs_with_bold(wr[j].paragraphs[0], txt)


def convert(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8")
    lines = text.split("\n")
    doc = Document()
    _style_base(doc)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # Bloc de code / diagramme ```
        if line.strip().startswith("```"):
            block: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            _add_code_block(doc, block)
            i += 1
            continue

        # Tableau Markdown (ligne | ... | suivie d'un séparateur ---)
        if line.lstrip().startswith("|") and i + 1 < n and re.match(r"\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            block = []
            while i < n and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_md_table(doc, block)
            doc.add_paragraph()
            continue

        # Séparateur horizontal
        if line.strip() == "---":
            i += 1
            continue

        # Titres
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            heading = doc.add_heading(level=level)
            _add_runs_with_bold(heading, m.group(2))
            i += 1
            continue

        # Citation >
        if line.startswith(">"):
            quote_lines = []
            while i < n and lines[i].startswith(">"):
                quote_lines.append(lines[i].lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            _add_runs_with_bold(p, " ".join(quote_lines))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = GREY
            continue

        # Liste à puces
        if re.match(r"^\s*[-*]\s+", line):
            content = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, content)
            i += 1
            continue

        # Liste numérotée
        if re.match(r"^\s*\d+\.\s+", line):
            content = re.sub(r"^\s*\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, content)
            i += 1
            continue

        # Ligne vide
        if not line.strip():
            i += 1
            continue

        # Paragraphe normal
        p = doc.add_paragraph()
        _add_runs_with_bold(p, line.strip())
        i += 1

    out_path = md_path.with_suffix(".docx")
    doc.save(str(out_path))
    return out_path


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
        "docs/certification/BLOC1_Collecter_Transformer_Securiser.md")
    out = convert(src)
    print(f"Word généré : {out}  ({out.stat().st_size:,} octets)")
